from pathlib import Path
import sys
import tempfile
from dataclasses import replace
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
import pandas as pd
from openpyxl import load_workbook
import batch_pre_sizing as bps
import gerar_casos_engstruct as eng
import gerar_casos_robustez_materia as robustez

expected = dict(cfg_d_min=20, cfg_d_max=100, cfg_bw_min=20,
                cfg_bw_max=50, cfg_h_min=5, cfg_h_max=15)
files = {'batch_pre_sizing_casos_materia.xlsx': 20,
         'robustez_casos_materia.xlsx': 4, 'engstruct_casos.xlsx': 320}
for name, count in files.items():
    old = load_workbook(ROOT / 'tmp/limites_casos/antes' / name)
    new = load_workbook(ROOT / name)
    assert old.sheetnames == new.sheetnames == ['Casos']
    before, after = old['Casos'], new['Casos']
    assert before.max_row == after.max_row == count + 1
    assert before.max_column == after.max_column
    headers = {c.column: c.value for c in after[1]}
    for row in before:
        for c in row:
            d = after[c.coordinate]
            assert (not c.has_style and not d.has_style) or c._style == d._style, (name, c.coordinate, 'style')
            if c.row > 1 and headers[c.column] in expected:
                assert d.value == expected[headers[c.column]], (name, c.coordinate)
            else:
                assert c.value == d.value, (name, c.coordinate, 'unrelated value')
    cases = bps.ler_planilha_casos(ROOT / name)
    assert len(cases) == count
    assert all(c.limites == bps.LimitesBusca() for c in cases)
    assert all(c.algoritmo.n_gen == 300 for c in cases)
    print(name, count, 'casos: limites corretos, outros valores e estilos preservados')
    old.close()
    new.close()

with tempfile.TemporaryDirectory(dir=ROOT / 'tmp/limites_casos') as tmp:
    tmp = Path(tmp)
    # Colunas ausentes/vazias usam os padrões; limites explícitos são respeitados.
    frame = pd.read_excel(ROOT / 'batch_pre_sizing_casos_materia.xlsx').head(1)
    frame = frame.drop(columns=['cfg_d_min'])
    frame['cfg_bw_max'] = float('nan')
    frame['cfg_h_max'] = 12.0
    frame.to_excel(tmp / 'fallback.xlsx', sheet_name='Casos', index=False)
    case = bps.ler_planilha_casos(tmp / 'fallback.xlsx')[0]
    assert case.limites == replace(bps.LimitesBusca(), h=(5.0, 12.0))
    print('Leitura: padrões ausentes/vazios e sobrescrita explícita corretos')

    # Executa os dois geradores reais em arquivos temporários, sem rodar o NSGA-II.
    with patch.object(sys, 'argv', ['gerar_casos_engstruct.py', '--saida', str(tmp / 'eng.xlsx')]):
        assert eng.main() == 0
    with patch.object(sys, 'argv', ['gerar_casos_robustez_materia.py', '--saida', str(tmp / 'robustez_casos_materia.xlsx')]):
        assert robustez.main() == 0
    for path, count in [(tmp / 'eng.xlsx', 320), (tmp / 'robustez_casos_materia.xlsx', 4)]:
        cases = bps.ler_planilha_casos(path)
        assert len(cases) == count
        assert all(c.limites == bps.LimitesBusca() for c in cases)
        print(path.name, 'gerador validado')
