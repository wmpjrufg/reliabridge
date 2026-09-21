r"""Atualiza apenas os limites de d, bw e h nas planilhas de entrada dos lotes.

Usa os padrões de LimitesBusca, preservando dados, espaçamentos, configuração do
algoritmo e formatação. Não executa otimização nem altera resultados históricos.

Uso:
    .venv\Scripts\python.exe atualizar_limites_casos.py
    .venv\Scripts\python.exe atualizar_limites_casos.py caminho.xlsx
"""

import argparse
from pathlib import Path

from openpyxl import load_workbook

from batch_pre_sizing import LimitesBusca, ler_planilha_casos

RAIZ = Path(__file__).resolve().parent
PLANILHAS = (
    "batch_pre_sizing_casos_materia.xlsx",
    "robustez_casos_materia.xlsx",
    "engstruct_casos.xlsx",
)


def atualizar(caminho: Path) -> int:
    """Valida a entrada e grava somente as seis colunas dimensionais existentes."""
    casos = ler_planilha_casos(caminho)
    limites = LimitesBusca()
    valores = {
        f"cfg_{nome}_{extremo}": getattr(limites, nome)[indice]
        for nome in ("d", "bw", "h")
        for indice, extremo in enumerate(("min", "max"))
    }
    workbook = load_workbook(caminho)
    try:
        sheet = workbook["Casos"]
        colunas = {cell.value: cell.column for cell in sheet[1]}
        ausentes = valores.keys() - colunas.keys()
        if ausentes:
            raise ValueError(f"{caminho.name}: colunas ausentes: {sorted(ausentes)}")
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, colunas["id"]).value is None:
                continue
            for nome, valor in valores.items():
                sheet.cell(row, colunas[nome], valor)
        workbook.save(caminho)
    finally:
        workbook.close()
    return len(casos)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("planilhas", nargs="*", type=Path)
    args = parser.parse_args()
    caminhos = args.planilhas or [RAIZ / nome for nome in PLANILHAS]
    for caminho in caminhos:
        total = atualizar(caminho)
        print(f"{caminho.name}: {total} casos; limites = {LimitesBusca()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
